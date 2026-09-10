"""
ExcelPredict - AI'nin Excel'deki veriyi disa aktarmadan tahmin etmesi.

Egitimli model her satirdaki metni okur, hangi konuya (intent) dustugunu
tahmin eder ve sonucu ayhni calisma sayfasina "Tahmin Edilen Konu" ve
"Guven (%)" sutunlarina yazar. Hicbir veri disari cikmaz.

Kullanim:
    python excel_predict.py dosya.xlsx
    python excel_predict.py dosya.xlsx --sutun C        # metin sutununu acikca belirt
    python excel_predict.py dosya.xlsx --sayfa Veri     # belirli bir sayfa
"""

import sys
import io
import argparse
import openpyxl
from openpyxl.utils import column_index_from_string

from brain import ChatBot

if sys.stdout.encoding.lower() not in ('utf-8', 'utf8'):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

MODEL_DIR = 'model'
TEXT_HEADERS = ('metin', 'soru', 'mesaj', 'girdi', 'cumle', 'yazi',
                'text', 'message', 'input', 'icerik')

PRED_COLUMNS = ('Tahmin Edilen Konu', 'Guven (%)')


def find_text_column(ws):
    """Ilk satirda metin oldugunu belirten baslik olan bir sutun bulur."""
    header = [c.value for c in ws[1]]
    if any(isinstance(h, str) and str(h).strip().lower() in TEXT_HEADERS for h in header):
        for i, h in enumerate(header):
            if isinstance(h, str) and str(h).strip().lower() in TEXT_HEADERS:
                return i + 1
    return 1


def main():
    parser = argparse.ArgumentParser(description='Excel icinde AI tahmini')
    parser.add_argument('file', help='.xlsx dosya yolu')
    parser.add_argument('--sutun', type=str, default='',
                        help='metin sutunu (ornek: B) - bos ise otomatik algilanir')
    parser.add_argument('--sayfa', type=str, default='',
                        help='calisma sayfasi adi (bos ise aktif sayfa)')
    args = parser.parse_args()

    bot = ChatBot()
    bot.load_model(MODEL_DIR)

    wb = openpyxl.load_workbook(args.file)
    ws = wb[args.sayfa] if args.sayfa else wb.active

    col = column_index_from_string(args.sutun.upper()) if args.sutun else find_text_column(ws)

    # Tahmin sutunlarini bul ya da mevcut son sutundan sonra ekle (veri silinmez)
    header = [c.value for c in ws[1]]
    pred_cols = {}
    next_free = len(header) + 1
    for name in PRED_COLUMNS:
        found = None
        for i, h in enumerate(header):
            if h == name:
                found = i + 1
                break
        if found is None:
            while next_free in pred_cols.values():
                next_free += 1
            found = next_free
            next_free += 1
            ws.cell(row=1, column=found, value=name)
        pred_cols[name] = found

    count = 0
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=col)
        text = cell.value
        if text is None or not str(text).strip():
            continue
        tag, conf = bot.predict(str(text))
        ws.cell(row=row, column=pred_cols['Tahmin Edilen Konu'], value=tag)
        ws.cell(row=row, column=pred_cols['Guven (%)'], value=conf)
        count += 1

    wb.save(args.file)
    print('=' * 50)
    print(f"TAMAM: {count} satir tahmin edildi -> {args.file}")
    print(f"Sonuclar ayni sayfaya yazildi: {PRED_COLUMNS}")


if __name__ == '__main__':
    main()